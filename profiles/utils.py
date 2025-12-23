import os
from io import BytesIO
from pdfminer.high_level import extract_text as extract_pdf_text
import docx

def extract_text_from_file(file_obj, file_extension):
    """
    Extracts text from PDF or DOCX file object.
    """
    text = ""
    try:
        if file_extension.lower() == '.pdf':
            # PDFMiner expects a path or file-like object
            text = extract_pdf_text(file_obj)
        elif file_extension.lower() in ['.docx', '.doc']:
            # python-docx requires a file-like object
            doc = docx.Document(file_obj)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            text = '\n'.join(full_text)
        else:
            raise ValueError("Unsupported file type")
        
        # Basic cleaning
        return clean_text(text)
    except Exception as e:
        print(f"Extraction error: {e}")
        return ""

def clean_text(text):
    """
    Removes extra whitespace and cleans up the text.
    """
    if not text:
        return ""
    
    # Split by lines, strip whitespace, remove empty lines
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    return '\n'.join(lines)

import requests
import re
from django.core.files.base import ContentFile
from PIL import Image

def download_and_process_icon(url, code_name):
    """
    Downloads custom icon from URL, resizes if raster, and returns ContentFile.
    """
    try:
        response = requests.get(url, timeout=10)
        response.raise_for_status()
        
        content_type = response.headers.get('Content-Type', '')
        
        if 'svg' in content_type or url.endswith('.svg'):
            # For SVG, we save as is (could add sanitization here)
            filename = f"{code_name}.svg"
            return ContentFile(response.content, name=filename), 'svg'
        
        # Assume raster image (png, jpg, etc)
        image = Image.open(BytesIO(response.content))
        
        # Resize to 64x64 max while keeping aspect ratio or just strict 64x64?
        # User said "Normalize size (64x64...)"
        image.thumbnail((64, 64), Image.Resampling.LANCZOS)
        
        output = BytesIO()
        image.save(output, format='PNG')
        filename = f"{code_name}.png"
        return ContentFile(output.getvalue(), name=filename), 'png'
        
    except Exception as e:
        print(f"Error downloading icon: {e}")
        return None, None

def extract_tech_from_bio(bio_text):
    """
    Extracts tech placeholders from bio text.
    Pattern: {{techn-name}}
    """
    pattern = r'\{\{([a-zA-Z0-9_-]+)\}\}'
    return re.findall(pattern, bio_text)

def get_absolute_media_url(file_field):
    """
    Returns absolute URL for a media file using BASE_URL from settings.
    
    Args:
        file_field: Django FileField or ImageField instance
        
    Returns:
        str: Absolute URL or None if file doesn't exist
    """
    if not file_field:
        return None
    
    from django.conf import settings
    
    # Get the relative URL
    relative_url = file_field.url
    
    # Combine with BASE_URL
    base_url = settings.BASE_URL.rstrip('/')
    return f"{base_url}{relative_url}"


def normalize_url(url: str) -> str:
    """
    Normalizes a URL by adding https:// if no protocol is present.
    
    Args:
        url: The URL string to normalize
        
    Returns:
        str: Normalized URL with protocol, or empty string if invalid
        
    Examples:
        - "github.com/user" -> "https://github.com/user"
        - "https://example.com" -> "https://example.com"
        - "user@example.com" -> "user@example.com"
        - "mailto:user@test.com" -> "mailto:user@test.com"
    """
    if not url or not isinstance(url, str):
        return ""
    
    url = url.strip()
    if not url:
        return ""
    
    # Already has protocol
    if url.startswith('http://') or url.startswith('https://'):
        return url
    
    # Email-related - don't add http
    if url.startswith('mailto:') or '@' in url:
        return url
    
    # Add https:// by default
    return f"https://{url}"

